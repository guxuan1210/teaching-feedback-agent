from io import BytesIO

from openpyxl import load_workbook


def test_export_contains_four_sheets_and_filtered_rows(client, mixed_feedback):
    response = client.get("/export.xlsx?date_from=2026-09-01&date_to=2026-09-30&class_id=C1")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["学生", "班级", "晚辅反馈", "专项反馈"]
    # C1 holds two daily records (李明 and 王芳); the special record lives in C2.
    assert book["晚辅反馈"].max_row == 3
    assert book["专项反馈"].max_row == 1


def test_student_export_includes_late_care_level(client):
    client.post("/catalog/students", data={
        "name": "档位学生",
        "grade": "四年级",
        "current_stage": "四阶",
        "late_care_level": "冲刺A",
    })

    book = load_workbook(BytesIO(client.get("/export.xlsx").content))
    sheet = book["学生"]
    assert [cell.value for cell in sheet[1]] == [
        "学生ID", "姓名", "年级", "九阶阶段", "晚辅档位", "状态"
    ]
    exported = next(row for row in sheet.iter_rows(values_only=True) if row[1] == "档位学生")
    assert exported[3:5] == ("四阶", "冲刺A")


def test_void_feedback_is_excluded_unless_requested(client, void_feedback):
    normal = load_workbook(BytesIO(client.get("/export.xlsx").content))
    included = load_workbook(BytesIO(client.get("/export.xlsx?status=all").content))
    assert normal["晚辅反馈"].max_row < included["晚辅反馈"].max_row
