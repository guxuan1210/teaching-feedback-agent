"""Filtered feedback workbook construction (openpyxl)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.catalog import repository as catalog_repository
from app.feedback.service import HistoryFilters, query_history


def _style_sheet(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(db: Session, filters: HistoryFilters) -> bytes:
    rows = query_history(db, filters)
    daily_rows = [r for r in rows if r.feedback_type == "daily"]
    special_rows = [r for r in rows if r.feedback_type == "special"]

    book = Workbook()

    ws_students = book.active
    ws_students.title = "学生"
    ws_students.append(["学生ID", "姓名", "年级", "九阶阶段", "状态"])
    for s in catalog_repository.list_students(db):
        ws_students.append([s.student_id, s.name, s.grade or "", s.current_stage or "", s.status])
    _style_sheet(ws_students, [22, 12, 10, 12, 10])

    ws_classes = book.create_sheet("班级")
    ws_classes.append(["班级ID", "名称", "年级", "类型", "状态"])
    for c in catalog_repository.list_classes(db):
        ws_classes.append([c.class_id, c.name, c.grade or "", c.class_type, c.status])
    _style_sheet(ws_classes, [22, 16, 10, 10, 10])

    ws_daily = book.create_sheet("晚辅反馈")
    ws_daily.append([
        "反馈ID", "日期", "时间", "班级", "学生", "老师",
        "知识掌握", "学习习惯", "心态与内驱力", "指标", "备注", "状态", "创建时间", "更新时间",
    ])
    for r in daily_rows:
        ws_daily.append([
            r.feedback_id, r.session_date, r.start_time, r.class_name, r.student_name, r.teacher_name,
            r.ratings.get("rating_knowledge"), r.ratings.get("rating_habit"), r.ratings.get("rating_mindset"),
            "；".join(r.indicator_texts), r.note or "", r.status, r.created_at, r.updated_at,
        ])
    _style_sheet(ws_daily, [20, 12, 8, 14, 10, 10, 10, 10, 12, 30, 20, 8, 24, 24])

    ws_special = book.create_sheet("专项反馈")
    ws_special.append([
        "反馈ID", "日期", "时间", "班级", "学生", "老师", "课程",
        "知识／技能", "课堂习惯", "指标", "备注", "状态", "创建时间", "更新时间",
    ])
    for r in special_rows:
        ws_special.append([
            r.feedback_id, r.session_date, r.start_time, r.class_name, r.student_name, r.teacher_name,
            r.course_name or "", r.ratings.get("rating_skill"), r.ratings.get("rating_habit"),
            "；".join(r.indicator_texts), r.note or "", r.status, r.created_at, r.updated_at,
        ])
    _style_sheet(ws_special, [20, 12, 8, 14, 10, 10, 14, 10, 10, 30, 20, 8, 24, 24])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
